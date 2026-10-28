"""将同一份报告快照导出为可交付的 DOCX 或 PDF。"""

from __future__ import annotations

import html
import re
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.textsplit import ALL_CANNOT_START
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont, TTFError
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from reportlab.platypus.paragraph import ParaLines, cjkU, makeCJKParaLine

from app.research.services.citation_format import REFERENCE_TITLES, clean_report_citations

_HEADING_RE = re.compile(r'^(#{1,6})\s+(.+?)\s*$')
_BULLET_RE = re.compile(r'^\s*[-*+]\s+(.+?)\s*$')
_NUMBERED_RE = re.compile(r'^\s*\d+[.)]\s+(.+?)\s*$')
_INLINE_LINK_RE = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')
_INLINE_BOLD_RE = re.compile(r'\*\*(.+?)\*\*')
_INLINE_ITALIC_RE = re.compile(r'(?<!\*)\*([^*]+)\*(?!\*)')
_CITATION_RE = re.compile(r'\[[1-9]\d*\](?!\()')
_SUBHEADING_RE = re.compile(r'^[一二三四五六七八九十百]+[、．]')


class CitationParagraph(Paragraph):
    """Keep a citation group with the preceding glyph in Chinese PDF wrapping.

    ReportLab's default CJK breaker splits individual glyphs, ignoring nobr.
    Retain its native fragment layout and superscript drawing, but pack the
    preceding character, citation group and closing punctuation as one unit.
    """

    def breakLinesCJK(self, maxWidths):
        if (getattr(self, '_splitpara', False) and hasattr(self, 'blPara')
                or self.bulletText or any(hasattr(f, 'cbDefn') or hasattr(f, 'lineBreak') for f in self.frags)
                or not any(getattr(f, 'rise', 0) > 0 for f in self.frags)):
            return super().breakLinesCJK(maxWidths)
        widths = list(maxWidths) if isinstance(maxWidths, (list, tuple)) else [maxWidths]
        self.height = 0
        units = []
        for fragment in self.frags:
            for character in fragment.text:
                glyph = cjkU(character, fragment, self.encoding)
                if units and (getattr(fragment, 'rise', 0) > 0 or character in ALL_CANNOT_START):
                    units[-1].append(glyph)
                else:
                    units.append([glyph])
        lines, glyphs, used = [], [], 0

        def finish():
            limit = widths[min(len(lines), len(widths) - 1)]
            lines.append(makeCJKParaLine(glyphs, limit, used, limit - used, False, True))

        for unit in units:
            size = sum(glyph.width for glyph in unit)
            limit = widths[min(len(lines), len(widths) - 1)]
            if glyphs and used + size > limit + 1e-6:
                finish()
                glyphs, used = [], 0
            glyphs.extend(unit)
            used += size
        if glyphs:
            finish()
        return ParaLines(kind=1, lines=lines)


def _pdf_font() -> str:
    name = 'PipeSightCJK'
    if name in pdfmetrics.getRegisteredFontNames():
        return name
    # Embed a local Unicode font when available so exported Chinese text does
    # not depend on the reader installing an Adobe CJK language pack.
    for filename in (
        '/System/Library/Fonts/Supplemental/Arial Unicode.ttf',
        '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc',
        '/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf',
        'C:/Windows/Fonts/simsun.ttc',
    ):
        if not Path(filename).is_file():
            continue
        try:
            pdfmetrics.registerFont(TTFont(name, filename))
            pdfmetrics.registerFontFamily(name, normal=name, bold=name, italic=name, boldItalic=name)
            return name
        except TTFError:
            continue
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    return 'STSong-Light'


def _has_reference_section(content: str) -> bool:
    return any(match and match.group(2).strip().casefold() in REFERENCE_TITLES
               for line in content.splitlines() if (match := _HEADING_RE.match(line.strip())))


def _docx_body(document: Any, text: str, *, reference: bool = False, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    text = text if reference else clean_report_citations(text)
    text = _INLINE_LINK_RE.sub(r'\1（\2）', text).replace('`', '').strip()
    if not reference and not style:
        paragraph.paragraph_format.first_line_indent = Pt(22)
        indent = paragraph._p.get_or_add_pPr().find(qn('w:ind'))
        indent.set(qn('w:firstLineChars'), '200')
    def add_text(value: str, *, bold=False, italic=False):
        offset = 0
        for match in ([] if reference else _CITATION_RE.finditer(value)):
            run = paragraph.add_run(value[offset:match.start()] + '\u2060')
            run.bold, run.italic = bold, italic
            run = paragraph.add_run(match[0])
            run.font.superscript = True
            run.bold, run.italic = bold, italic
            paragraph.add_run('\u2060')
            offset = match.end()
        run = paragraph.add_run(value[offset:])
        run.bold, run.italic = bold, italic

    offset = 0
    for match in re.finditer(r'\*\*(.+?)\*\*|(?<!\*)\*([^*]+)\*(?!\*)', text):
        add_text(text[offset:match.start()])
        add_text(match[1] or match[2], bold=match[1] is not None, italic=match[2] is not None)
        offset = match.end()
    add_text(text[offset:])
    if reference:
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(4)
        for run in paragraph.runs:
            run.font.size = Pt(9.5)


def safe_file_stem(title: str, *, is_draft: bool) -> str:
    stem = re.sub(r'[\\/:*?"<>|\r\n]+', '_', str(title or '').strip()).strip(' ._') or '科研报告'
    return f'草稿_{stem}' if is_draft else stem


def _plain_inline(value: str) -> str:
    value = _INLINE_LINK_RE.sub(r'\1（\2）', value)
    value = value.replace('`', '')
    return _INLINE_BOLD_RE.sub(r'\1', _INLINE_ITALIC_RE.sub(r'\1', value)).strip()


def _report_body(report: Any) -> str:
    content = str(getattr(report, 'content', '') or '').strip()
    if content:
        return content
    sections = getattr(report, 'sections', None) or []
    return '\n\n'.join(
        f'{"#" * max(1, min(int(section.get("level") or 2), 6))} {section.get("title", "")}\n\n'
        f'{section.get("content", "")}'
        for section in sections
        if isinstance(section, dict) and str(section.get('title') or '').strip()
    )


def _citation_lines(report: Any) -> list[str]:
    citations = getattr(report, 'citations', None) or []
    lines: list[str] = []
    for index, citation in enumerate(citations, start=1):
        if not isinstance(citation, dict):
            lines.append(f'{index}. {citation}')
            continue
        title = str(citation.get('title') or citation.get('name') or citation.get('id') or '未命名资料')
        source = str(citation.get('source') or citation.get('sourceName') or '').strip()
        url = str(citation.get('url') or '').strip()
        suffix = '；'.join(part for part in (source, url) if part)
        lines.append(f'{index}. {title}{f"（{suffix}）" if suffix else ""}')
    return lines


def _metadata_line(report: Any) -> str:
    owner = str(getattr(report, 'owner_name', '') or '').strip()
    report_type = {
        'research-demand': '科研需求报告',
        'technology-route': '技术路线报告',
        'project-demand': '科技项目需求填报表',
        'daily': '科技情报日报',
        'weekly': '科技情报周报',
        'monthly': '科技情报月报',
    }.get(str(getattr(report, 'report_type', '') or ''), '')
    metadata = dict(getattr(report, 'metadata', None) or {})
    parts = [f'类型：{report_type}'] if report_type else []
    if owner:
        parts.append(f'负责人：{owner}')
    period = metadata.get('period')
    if isinstance(period, (list, tuple)) and period:
        parts.append(f'周期：{" 至 ".join(str(item) for item in period)}')
    return '　'.join(parts)


def _docx_page_number(document: Any) -> None:
    section = document.sections[0]
    section.footer_distance = Inches(0.35)
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    run = paragraph.add_run('1')
    run.font.name = 'Arial'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(102, 102, 102)
    field.append(run._r)
    paragraph._p.append(field)
    update = document.settings.element.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields')
        document.settings.element.append(update)
    update.set(qn('w:val'), 'true')


def build_docx(report: Any, *, is_draft: bool) -> tuple[str, str, bytes]:
    """返回文件名、媒体类型和 DOCX 二进制内容。"""
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    _docx_page_number(document)
    normal = document.styles['Normal']
    normal.font.name = 'Songti SC'
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Songti SC')
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4', 'Heading 5', 'Heading 6']:
        style = document.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.name = 'Songti SC'
        fonts = style.element.get_or_add_rPr().rFonts
        fonts.set(qn('w:eastAsia'), 'Songti SC')
        for attribute in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            fonts.attrib.pop(qn(f'w:{attribute}'), None)
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
        if name.startswith('Heading '):
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True

    title = str(getattr(report, 'title', '') or '科研报告')
    document.core_properties.title = title
    document.core_properties.author = str(getattr(report, 'owner_name', '') or 'PipeSight')
    document.styles['Title'].font.size = Pt(20)
    title_paragraph = document.add_heading(title, 0)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata_text = _metadata_line(report)
    if metadata_text:
        metadata = document.add_paragraph(metadata_text)
        metadata.style = document.styles['Subtitle']

    content = _report_body(report)
    reference = False
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = _HEADING_RE.match(line)
        if heading:
            reference = heading.group(2).strip().casefold() in REFERENCE_TITLES
            if len(heading.group(1)) == 1 and heading.group(2).strip() == title:
                continue
            document.add_heading(_plain_inline(heading.group(2)), min(max(len(heading.group(1)) - 1, 1), 6))
            continue
        if not reference and _SUBHEADING_RE.match(line):
            document.add_heading(_plain_inline(line), 3)
            continue
        bullet = _BULLET_RE.match(line)
        if bullet:
            _docx_body(document, bullet.group(1), reference=reference, style='List Bullet')
            continue
        _docx_body(document, line, reference=reference)

    citation_lines = _citation_lines(report)
    if citation_lines and not _has_reference_section(content):
        document.add_heading('参考资料', 1)
        for line in citation_lines:
            _docx_body(document, line, reference=True)

    output = BytesIO()
    document.save(output)
    return (
        f'{safe_file_stem(title, is_draft=is_draft)}.docx',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        output.getvalue(),
    )


def _pdf_inline(value: str, *, citation: bool = False) -> str:
    escaped = html.escape(value, quote=False)
    escaped = _INLINE_LINK_RE.sub(r'\1（\2）', escaped)
    escaped = _INLINE_BOLD_RE.sub(r'<b>\1</b>', escaped)
    escaped = _INLINE_ITALIC_RE.sub(r'<i>\1</i>', escaped)
    return _CITATION_RE.sub(r'<super>\g<0></super>', escaped) if citation else escaped


def build_pdf(report: Any, *, is_draft: bool) -> tuple[str, str, bytes]:
    """返回文件名、媒体类型和 PDF 二进制内容，使用内置 CJK 字体避免中文变方框。"""
    font_name = _pdf_font()
    title = str(getattr(report, 'title', '') or '科研报告')
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'PipeSightTitle', parent=styles['Title'], fontName=font_name, fontSize=20, leading=26, alignment=TA_CENTER
    )
    metadata_style = ParagraphStyle(
        'PipeSightMetadata',
        parent=styles['Normal'],
        fontName=font_name,
        fontSize=9,
        leading=14,
        textColor='#666666',
    )
    heading_styles = {
        level: ParagraphStyle(
            f'PipeSightHeading{level}',
            parent=styles['Heading1'],
            fontName=font_name,
            fontSize=max(11, 18 - level),
            leading=max(15, 23 - level),
            spaceBefore=8,
            spaceAfter=4,
            keepWithNext=True,
        )
        for level in range(1, 7)
    }
    body_style = ParagraphStyle(
        'PipeSightBody', parent=styles['BodyText'], fontName=font_name, fontSize=10.5, leading=18, spaceAfter=6,
        firstLineIndent=21, wordWrap='CJK'
    )
    reference_style = ParagraphStyle('PipeSightReference', parent=body_style, firstLineIndent=0,
                                     fontSize=9.5, leading=14, spaceAfter=4)
    bullet_style = ParagraphStyle(
        'PipeSightBullet', parent=body_style, leftIndent=14, firstLineIndent=-10, bulletIndent=0
    )

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=title,
        author=str(getattr(report, 'owner_name', '') or 'PipeSight'),
    )
    story: list[Any] = [Paragraph(_pdf_inline(title), title_style), Spacer(1, 5 * mm)]
    metadata_text = _metadata_line(report)
    if metadata_text:
        story.append(Paragraph(_pdf_inline(metadata_text), metadata_style))
        story.append(Spacer(1, 5 * mm))

    content = _report_body(report)
    reference = False
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            # Semantic paragraphs already have spaceAfter. Browser-edited content
            # may contain several blank lines; don't multiply them into layout gaps.
            continue
        heading = _HEADING_RE.match(line)
        if heading:
            reference = heading.group(2).strip().casefold() in REFERENCE_TITLES
            if len(heading.group(1)) == 1 and heading.group(2).strip() == title:
                continue
            level = min(len(heading.group(1)), 6)
            story.append(Paragraph(_pdf_inline(heading.group(2)), heading_styles[level]))
            continue
        if not reference and _SUBHEADING_RE.match(line):
            story.append(Paragraph(_pdf_inline(line), heading_styles[3]))
            continue
        line = line if reference else clean_report_citations(line)
        bullet = _BULLET_RE.match(line)
        if bullet:
            story.append(Paragraph(_pdf_inline(bullet.group(1), citation=not reference), bullet_style, bulletText='•'))
            continue
        story.append(CitationParagraph(_pdf_inline(line, citation=not reference), reference_style if reference else body_style))

    citation_lines = _citation_lines(report)
    if citation_lines and not _has_reference_section(content):
        story.append(Paragraph('参考资料', heading_styles[1]))
        story.extend(Paragraph(_pdf_inline(line), reference_style) for line in citation_lines)

    def page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont(font_name, 9)
        canvas.setFillColorRGB(0.4, 0.4, 0.4)
        canvas.drawCentredString(doc.pagesize[0] / 2, 9 * mm, str(doc.page))
        canvas.restoreState()

    document.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return f'{safe_file_stem(title, is_draft=is_draft)}.pdf', 'application/pdf', output.getvalue()


def build_file(report: Any, file_format: str, *, is_draft: bool) -> tuple[str, str, bytes]:
    normalized = str(file_format or '').strip().lower()
    if normalized == 'docx':
        return build_docx(report, is_draft=is_draft)
    if normalized == 'pdf':
        return build_pdf(report, is_draft=is_draft)
    raise ValueError(f'不支持的导出格式：{file_format}')


__all__ = ['build_docx', 'build_file', 'build_pdf', 'safe_file_stem']
