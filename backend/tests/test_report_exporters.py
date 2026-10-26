from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas

from app.research.services.exporters import build_docx, build_pdf, safe_file_stem

MINIMUM_PDF_BYTES = 1000


def _report() -> SimpleNamespace:
    return SimpleNamespace(
        title='管道完整性/年度报告',
        owner_name='测试用户',
        version_no=3,
        content='# 研究结论\n\n正文内容 **需要复核**。\n\n- 第一条建议',
        sections=[{'title': '研究结论', 'level': 1, 'content': '正文内容'}],
        citations=[{'title': '公开资料', 'source': '管网来源', 'url': 'https://example.com'}],
    )


def test_docx_export_contains_report_snapshot() -> None:
    file_name, media_type, content = build_docx(_report(), is_draft=True)

    assert file_name == '草稿_管道完整性_年度报告.docx'
    assert media_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert content.startswith(b'PK')
    with ZipFile(BytesIO(content)) as archive:
        document_xml = archive.read('word/document.xml').decode('utf-8')
    assert '研究结论' in document_xml
    assert '正文内容' in document_xml
    assert '公开资料' in document_xml


def test_pdf_export_is_real_pdf_and_uses_safe_filename() -> None:
    file_name, media_type, content = build_pdf(_report(), is_draft=False)

    assert file_name == '管道完整性_年度报告.pdf'
    assert media_type == 'application/pdf'
    assert content.startswith(b'%PDF-')
    assert len(content) > MINIMUM_PDF_BYTES


def test_export_filename_removes_path_separators() -> None:
    assert safe_file_stem('../draft/report', is_draft=False) == 'draft_report'


def test_export_body_citations_indent_superscript_and_no_duplicate_reference_section() -> None:
    report = _report()
    report.content = ('# ' + report.title + '\n\n## 研究内容\n\n一、研究背景\n\n'
                      '[2]新闻正文报道，系统已上线[2]（新闻正文，2025-02-09）。\n\n'
                      '## 参考资料\n\n[2] 新闻来源；2025-02-09；https://example.org')
    _, _, data = build_docx(report, is_draft=False)
    doc = Document(BytesIO(data))
    paragraphs = doc.paragraphs
    body = next(p for p in paragraphs if '系统已上线' in p.text)
    assert body.text.replace('\u2060', '') == '新闻正文报道，系统已上线[2]。'
    assert '\u2060[2]\u2060' in body.text
    assert any(run.text == '[2]' and run.font.superscript for run in body.runs)
    indent = body._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ind')
    assert indent is not None and indent.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}firstLineChars') == '200'
    reference = next(p for p in paragraphs if '新闻来源' in p.text)
    assert all(not run.font.superscript for run in reference.runs)
    assert sum(p.text == '参考资料' for p in paragraphs) == 1
    assert sum(p.text == report.title for p in paragraphs) == 1


def test_pdf_citation_markup_and_indent_are_real_layout_instructions(monkeypatch) -> None:
    from app.research.services import exporters
    calls = []
    original = exporters.Paragraph
    def capture(text, style, *args, **kwargs):
        calls.append((text, style))
        return original(text, style, *args, **kwargs)
    monkeypatch.setattr(exporters, 'Paragraph', capture)
    monkeypatch.setattr(exporters, 'CitationParagraph', capture)
    report = _report()
    report.content = '## 研究内容\n\n[2]新闻正文报道，系统已上线[2]（新闻正文，2025-02-09）。\n\n## 参考资料\n\n[2] 新闻；2025-02-09'
    build_pdf(report, is_draft=False)
    text, style = next(item for item in calls if '系统已上线' in item[0])
    assert text == '新闻正文报道，系统已上线<super>[2]</super>。'
    assert style.firstLineIndent == 2 * style.fontSize
    reference, ref_style = next(item for item in calls if '新闻；' in item[0])
    assert '<super>' not in reference and ref_style.firstLineIndent == 0


def test_cjk_pdf_citations_cannot_be_stranded_on_a_line_without_their_claim() -> None:
    from app.research.services import exporters
    font = exporters._pdf_font()
    paragraph_class = getattr(exporters, 'CitationParagraph', exporters.Paragraph)
    for width in range(42, 180, 3):
        paragraph = paragraph_class('检测方法需要验证<super>[1][2]</super>。后续继续研究。',
                                    ParagraphStyle('CitationWrap', fontName=font, fontSize=10.5,
                                                   leading=18, wordWrap='CJK'))
        paragraph.wrap(width, 800)
        lines = [''.join(fragment.text for fragment in line.words) for line in paragraph.blPara.lines]
        assert all(not line.startswith(('[', ']')) for line in lines), (width, lines)
        assert ''.join(lines) == '检测方法需要验证[1][2]。后续继续研究。'


def test_word_footer_has_a_centered_native_page_field_not_a_fixed_number() -> None:
    _, _, content = build_docx(_report(), is_draft=False)
    doc = Document(BytesIO(content))
    footer = doc.sections[0].footer.paragraphs[0]
    assert footer.alignment == WD_ALIGN_PARAGRAPH.CENTER
    fields = footer._p.findall('.//' + qn('w:fldSimple'))
    assert len(fields) == 1 and fields[0].get(qn('w:instr')).strip() == 'PAGE'
    assert doc.settings.element.find(qn('w:updateFields')).get(qn('w:val')) == 'true'


def test_pdf_footer_numbers_every_page_consecutively_at_bottom_center(monkeypatch) -> None:
    drawn = []
    original = canvas.Canvas.drawCentredString

    def capture(self, x, y, label, *args, **kwargs):
        drawn.append((x, y, label))
        return original(self, x, y, label, *args, **kwargs)

    monkeypatch.setattr(canvas.Canvas, 'drawCentredString', capture)
    report = _report()
    report.content = '## 正文\n\n' + ('页面内容用于确认页码连续且不遮挡正文。\n\n' * 120)
    build_pdf(report, is_draft=False)
    numbers = [(x, y, label) for x, y, label in drawn if str(label).isdigit()]
    assert len(numbers) > 1
    assert [label for _, _, label in numbers] == [str(n) for n in range(1, len(numbers) + 1)]
    assert all(x == A4[0] / 2 and 0 < y < 18 * 72 / 25.4 for x, y, _ in numbers)
